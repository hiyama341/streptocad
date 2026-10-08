import io
import json
import zipfile

import pandas as pd
import pytest
from Bio.Seq import Seq
from Bio.SeqRecord import SeqRecord

from streptocad.output_packaging import (
    PROTOCOLS,
    ROLES,
    SHELF_ORDER,
    OutputPackage,
    environment_record,
    run_id,
    sanitize_path_component,
)

TIMESTAMP = "20261007-121503Z"


def record(name, id_=None, seq="ATGC"):
    rec = SeqRecord(Seq(seq * 10), id=id_ or name, name=name, description="")
    rec.annotations["molecule_type"] = "DNA"
    return rec


def names_in(payload):
    with zipfile.ZipFile(io.BytesIO(payload)) as zf:
        return zf.namelist()


def read_in(payload, suffix):
    """Read the one member whose path ends with ``suffix``."""
    with zipfile.ZipFile(io.BytesIO(payload)) as zf:
        match = [n for n in zf.namelist() if n.endswith(suffix)]
        assert len(match) == 1, f"expected one {suffix}, got {match}"
        return zf.read(match[0]).decode("utf-8")


def strip_root(payload):
    """Member paths with the single root directory removed."""
    out = []
    for name in names_in(payload):
        head, _, tail = name.partition("/")
        assert head.startswith("streptocad_"), f"{name} is not under one root"
        out.append(tail)
    return out


# -- naming ------------------------------------------------------------------


@pytest.mark.parametrize(
    "raw, expected",
    [
        ("pOEX-PKasO", "pOEX-PKasO"),
        ("pOEX/kasO", "pOEX_kasO"),
        ("pCas9_SCO5892(4056120)", "pCas9_SCO5892_4056120"),
        ("sgRNA #1", "sgRNA_1"),
        ("../../etc/passwd", "etc_passwd"),
        ("...", "unnamed"),
    ],
)
def test_sanitize_path_component(raw, expected):
    assert sanitize_path_component(raw) == expected


def test_run_id_is_filesystem_safe_and_sortable():
    name = run_id("w1", TIMESTAMP)
    assert name == "streptocad_w1-overexpression_20261007-121503Z"
    # ':' is illegal on NTFS, which is what broke extraction on Windows.
    assert ":" not in name


def test_environment_record_names_the_key_packages():
    env = environment_record()
    assert "python" in env
    assert "pandas" in env


# -- structure ---------------------------------------------------------------


def test_everything_lands_under_one_root_directory():
    payload = OutputPackage(
        "w1", [{"role": "primer.pcr", "content": pd.DataFrame({"a": [1]})}]
    ).to_zip_bytes(TIMESTAMP)

    roots = {n.split("/")[0] for n in names_in(payload)}
    assert roots == {"streptocad_w1-overexpression_20261007-121503Z"}


def test_artifacts_land_on_their_declared_shelf():
    payload = OutputPackage(
        "w2",
        outputs=[
            {"role": "primer.order_idt", "content": pd.DataFrame({"Name": ["p"]})},
            {"role": "sgrna.selected", "content": pd.DataFrame({"g": ["ATGC"]})},
            {"role": "plasmid", "content": [record("pCRISPR_SCO5892")]},
            {"role": "analysis.primer_qc", "content": pd.DataFrame({"hp": [1]})},
        ],
        inputs=[{"role": "input.genome", "content": record("genome")}],
    ).to_zip_bytes(TIMESTAMP)

    paths = strip_root(payload)
    assert "2_primers/oligo_order_idt.csv" in paths
    assert "3_sgrnas/sgrna_selected.csv" in paths
    assert "1_plasmids/01_pCRISPR_SCO5892.gb" in paths
    assert "4_analysis/primer_qc_hairpins.csv" in paths
    assert "6_inputs/input_genome.gb" in paths
    assert "6_inputs/run_parameters.json" in paths
    assert "4_analysis/environment.json" in paths


def test_no_numeric_prefixes_on_data_files():
    payload = OutputPackage(
        "w3",
        outputs=[
            {"role": "primer.pcr", "content": pd.DataFrame({"a": [1]})},
            {"role": "primer.order_idt", "content": pd.DataFrame({"a": [1]})},
            {"role": "sgrna.all", "content": pd.DataFrame({"a": [1]})},
            {"role": "analysis.golden_gate_overhangs", "content": pd.DataFrame({"a": [1]})},
        ],
    ).to_zip_bytes(TIMESTAMP)

    data_files = [
        p.split("/", 1)[1]
        for p in strip_root(payload)
        if p.endswith(".csv") and "/" in p
    ]
    assert data_files, "expected some CSVs"
    for name in data_files:
        assert not name[:2].isdigit(), f"{name} still carries a numeric prefix"


def test_unused_shelf_is_omitted_and_explained_in_the_readme():
    payload = OutputPackage(
        "w1",
        outputs=[
            {"role": "plasmid", "content": [record("pOEX_SCO5087")]},
            {"role": "primer.pcr", "content": pd.DataFrame({"a": [1]})},
        ],
    ).to_zip_bytes(TIMESTAMP)

    paths = strip_root(payload)
    assert not any(p.startswith("3_sgrnas/") for p in paths)

    readme = read_in(payload, "00_READ_ME_FIRST.md")
    assert "3_sgrnas/ -- not used" in readme
    assert "uses no guide RNAs" in readme


def test_readme_ships_as_html_and_markdown_and_names_the_next_step():
    payload = OutputPackage(
        "w1", [{"role": "primer.order_idt", "content": pd.DataFrame({"a": [1]})}]
    ).to_zip_bytes(TIMESTAMP)

    paths = strip_root(payload)
    assert "00_READ_ME_FIRST.md" in paths
    assert "00_READ_ME_FIRST.html" in paths

    readme = read_in(payload, "00_READ_ME_FIRST.md")
    assert "idtdna.com" in readme
    assert "oligo_order_idt" in readme


# -- plasmids ----------------------------------------------------------------


def test_plasmids_are_numbered_leading_and_zero_padded():
    plasmids = [record(f"pOEX_SCO{5087 + i}") for i in range(12)]
    payload = OutputPackage(
        "w1", [{"role": "plasmid", "content": plasmids}]
    ).to_zip_bytes(TIMESTAMP)

    gb = sorted(p for p in strip_root(payload) if p.endswith(".gb"))
    assert len(gb) == 12
    assert gb[0].startswith("1_plasmids/01_")
    assert gb[-1].startswith("1_plasmids/12_")
    # Lexicographic order is design order, which the old trailing "_0".."_19"
    # suffix did not give.
    assert gb == sorted(gb)


def test_plasmid_filenames_cannot_escape_their_shelf():
    payload = OutputPackage(
        "w1", [{"role": "plasmid", "content": [record("pOEX/kasO")]}]
    ).to_zip_bytes(TIMESTAMP)

    gb = [p for p in strip_root(payload) if p.endswith(".gb")]
    assert gb == ["1_plasmids/01_pOEX_kasO.gb"]


def test_a_record_name_with_a_space_still_produces_an_archive():
    """SeqIO rejects whitespace in a LOCUS name, which used to kill the download."""
    payload = OutputPackage(
        "w1", [{"role": "plasmid", "content": [record("pOEX kasO 1")]}]
    ).to_zip_bytes(TIMESTAMP)

    gb = [p for p in strip_root(payload) if p.endswith(".gb")]
    assert len(gb) == 1
    assert " " not in gb[0]


def test_an_empty_plasmid_list_warns_instead_of_shipping_silence():
    package = OutputPackage("w1", [{"role": "plasmid", "content": []}])
    payload = package.to_zip_bytes(TIMESTAMP)

    paths = strip_root(payload)
    assert "1_plasmids/NO_PLASMIDS_ASSEMBLED.txt" in paths
    assert any("No plasmids were assembled" in w for w in package.warnings)
    assert "No plasmids were assembled" in read_in(payload, "00_READ_ME_FIRST.md")


# -- IDT plate ---------------------------------------------------------------


def test_idt_plate_ships_as_xlsx_with_one_worksheet_per_plate():
    from openpyxl import load_workbook

    from streptocad.primers.idt_plates import assign_plate_wells

    plates = assign_plate_wells(
        [f"oligo_{i}" for i in range(100)], ["ATGC"] * 100
    )
    payload = OutputPackage(
        "w1", [{"role": "primer.order_idt_plate", "content": plates}]
    ).to_zip_bytes(TIMESTAMP)

    assert "2_primers/oligo_order_idt_plate.xlsx" in strip_root(payload)

    with zipfile.ZipFile(io.BytesIO(payload)) as zf:
        name = [n for n in zf.namelist() if n.endswith(".xlsx")][0]
        workbook = load_workbook(io.BytesIO(zf.read(name)))

    assert workbook.sheetnames == ["Plate1", "Plate2"]
    assert [c.value for c in workbook["Plate1"][1]] == [
        "Well Position",
        "Name",
        "Sequence",
    ]


# -- protocols ---------------------------------------------------------------


def test_protocols_ship_as_html_and_markdown_under_canonical_names():
    payload = OutputPackage(
        "w1",
        outputs=[{"role": "primer.pcr", "content": pd.DataFrame({"a": [1]})}],
        protocols=["conjugation", "troubleshooting"],
    ).to_zip_bytes(TIMESTAMP)

    paths = strip_root(payload)
    # The source files are misspelled on disk; the shipped names are not.
    assert "5_protocols/conjugation_protocol.html" in paths
    assert "5_protocols/conjugation_protocol.md" in paths
    assert "5_protocols/troubleshooting_tips.html" in paths
    assert not any("protcol" in p for p in paths)
    assert not any("trouble_shooting" in p for p in paths)


def test_every_protocol_key_resolves_to_a_file_on_disk():
    """Guards against a protocol being renamed on disk without the table."""
    package = OutputPackage(
        "w1",
        outputs=[{"role": "primer.pcr", "content": pd.DataFrame({"a": [1]})}],
        protocols=list(PROTOCOLS),
    )
    payload = package.to_zip_bytes(TIMESTAMP)

    assert package.warnings == [], f"unreadable protocols: {package.warnings}"
    html = [p for p in strip_root(payload) if p.endswith(".html")]
    # One per protocol, plus the README.
    assert len(html) == len(PROTOCOLS) + 1


def test_a_missing_protocol_warns_but_still_ships_the_designs(tmp_path):
    package = OutputPackage(
        "w1",
        outputs=[{"role": "plasmid", "content": [record("pOEX_SCO5087")]}],
        protocols=["conjugation"],
        protocols_dir=str(tmp_path),
    )
    payload = package.to_zip_bytes(TIMESTAMP)

    assert any("could not be read" in w for w in package.warnings)
    assert any(p.endswith(".gb") for p in strip_root(payload))


# -- contract --------------------------------------------------------------


def test_an_unknown_role_raises_rather_than_misfiling():
    with pytest.raises(KeyError, match="unknown output role"):
        OutputPackage("w1", [{"role": "primer.typo", "content": pd.DataFrame()}])


def test_an_unknown_protocol_key_raises():
    with pytest.raises(KeyError, match="unknown protocol key"):
        OutputPackage("w1", [], protocols=["not_a_protocol"])


def test_a_table_role_rejects_a_non_dataframe():
    package = OutputPackage("w1", [{"role": "primer.pcr", "content": "not a frame"}])
    with pytest.raises(TypeError, match="must be a.*DataFrame"):
        package.to_zip_bytes(TIMESTAMP)


def test_parameters_round_trip_as_json():
    params = {"genes": ["SCO5087"], "gc_upper": 0.8}
    payload = OutputPackage("w1", [], parameters=params).to_zip_bytes(TIMESTAMP)

    assert json.loads(read_in(payload, "run_parameters.json")) == params


def test_every_role_declares_a_known_shelf_and_a_blurb():
    for role, (shelf, _filename, kind, blurb) in ROLES.items():
        assert shelf in SHELF_ORDER, f"{role} points at unknown shelf {shelf}"
        assert kind in {"table", "records", "plate", "text", "json"}, role
        assert blurb and blurb[0].isupper(), f"{role} needs a readable blurb"


def test_packaging_twice_does_not_accumulate():
    """The old buffer was never truncated, so a second call returned the first."""
    package = OutputPackage(
        "w1", [{"role": "primer.pcr", "content": pd.DataFrame({"a": [1]})}]
    )
    first = package.to_zip_bytes(TIMESTAMP)
    second = package.to_zip_bytes(TIMESTAMP)

    assert sorted(names_in(first)) == sorted(names_in(second))


def test_packaging_does_not_mutate_the_callers_records():
    """A record may be declared under more than one role."""
    plasmid = record("a_very_long_plasmid_name_beyond_locus")
    original_name = plasmid.name

    OutputPackage(
        "w1",
        [
            {"role": "plasmid", "content": [plasmid]},
            {"role": "plasmid.all", "content": [plasmid]},
        ],
    ).to_zip_bytes(TIMESTAMP)

    assert plasmid.name == original_name


def test_readme_lists_every_shipped_file_with_its_own_blurb():
    """The plate file went missing from the listing, and inputs showed the
    protocol fallback blurb, because neither recorded its role."""
    from streptocad.primers.idt_plates import assign_plate_wells

    package = OutputPackage(
        "w1",
        outputs=[
            {
                "role": "primer.order_idt_plate",
                "content": assign_plate_wells(["o1"], ["ATGC"]),
            }
        ],
        inputs=[{"role": "input.sequences", "content": [record("gene_a")]}],
    )
    payload = package.to_zip_bytes(TIMESTAMP)
    readme = read_in(payload, "00_READ_ME_FIRST.md")

    assert "oligo_order_idt_plate.xlsx" in readme
    assert ROLES["primer.order_idt_plate"][3] in readme
    assert ROLES["input.sequences"][3] in readme
    assert "Bench instructions" not in readme

    # Every non-plasmid shipped file is named in the README.
    for path, role in package.written.items():
        if role == "plasmid" and path.endswith(".gb"):
            continue
        assert path.split("/", 1)[-1] in readme, f"{path} missing from README"


# -- the plate must cover the whole order ------------------------------------


def test_an_incomplete_plate_order_is_flagged():
    """The plate and the tube sheet are one order in two formats."""
    from streptocad.primers.idt_plates import assign_plate_wells

    tube = pd.DataFrame({"Name": [f"o{i}" for i in range(10)], "Sequence": ["ATGC"] * 10})
    partial = assign_plate_wells([f"o{i}" for i in range(4)], ["ATGC"] * 4)

    package = OutputPackage(
        "w5",
        [
            {"role": "primer.order_idt", "content": tube},
            {"role": "primer.order_idt_plate", "content": partial},
        ],
    )
    payload = package.to_zip_bytes(TIMESTAMP)

    assert any("lists 4 oligos but the tube order lists 10" in w for w in package.warnings)
    assert "Order from oligo_order_idt.csv" in read_in(payload, "00_READ_ME_FIRST.md")


def test_a_complete_plate_order_is_not_flagged():
    from streptocad.primers.idt_plates import idt_order_df_to_idt_plates

    tube = pd.DataFrame({"Name": [f"o{i}" for i in range(10)], "Sequence": ["ATGC"] * 10})

    package = OutputPackage(
        "w5",
        [
            {"role": "primer.order_idt", "content": tube},
            {"role": "primer.order_idt_plate", "content": idt_order_df_to_idt_plates(tube)},
        ],
    )
    package.to_zip_bytes(TIMESTAMP)

    assert package.warnings == []
