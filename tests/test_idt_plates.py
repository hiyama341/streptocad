import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from streptocad.primers.idt_plates import (
    IDT_PLATE_COLUMNS,
    assign_plate_wells,
    idt_order_df_to_idt_plates,
    idt_plates_to_xlsx_bytes,
    sanitize_idt_oligo_name,
    well_positions,
)



def test_well_positions_fills_row_major():
    wells = well_positions()
    assert len(wells) == 96
    assert wells[:13] == [
        "A1",
        "A2",
        "A3",
        "A4",
        "A5",
        "A6",
        "A7",
        "A8",
        "A9",
        "A10",
        "A11",
        "A12",
        "B1",
    ]
    assert wells[-1] == "H12"


def test_well_positions_rejects_impossible_geometry():
    with pytest.raises(ValueError):
        well_positions(rows=27)
    with pytest.raises(ValueError):
        well_positions(cols=0)


def test_assign_plate_wells_uses_idt_headers_in_order():
    plates = assign_plate_wells(["oligo_a", "oligo_b"], ["atgc", "ggcc"])
    assert len(plates) == 1
    assert list(plates[0].columns) == IDT_PLATE_COLUMNS
    assert plates[0]["Well Position"].tolist() == ["A1", "A2"]
    # Sequences are upper-cased; IDT's templates are upper case throughout.
    assert plates[0]["Sequence"].tolist() == ["ATGC", "GGCC"]


def test_assign_plate_wells_splits_at_96_and_restarts_at_a1():
    names = [f"oligo_{i}" for i in range(100)]
    seqs = ["ATGC"] * 100

    plates = assign_plate_wells(names, seqs)

    assert len(plates) == 2
    assert len(plates[0]) == 96
    assert len(plates[1]) == 4
    assert plates[0]["Well Position"].iloc[-1] == "H12"
    assert plates[1]["Well Position"].tolist() == ["A1", "A2", "A3", "A4"]
    assert plates[1]["Name"].tolist() == [
        "oligo_96",
        "oligo_97",
        "oligo_98",
        "oligo_99",
    ]


def test_assign_plate_wells_empty_input_gives_no_plates():
    assert assign_plate_wells([], []) == []


def test_assign_plate_wells_rejects_mismatched_lengths():
    with pytest.raises(ValueError):
        assign_plate_wells(["a", "b"], ["ATGC"])


@pytest.mark.parametrize(
    "raw, expected",
    [
        ("primer_fwd_SCO5087", "primer_fwd_SCO5087"),
        ("pCas9_SCO5892(4056120)", "pCas9_SCO5892_4056120"),
        ("sgRNA_#1_SCO5892", "sgRNA_1_SCO5892"),
        ("primer fwd SCO5087", "primer_fwd_SCO5087"),
        ("pOEX/kasO", "pOEX_kasO"),
        ("CW1060_LLPMBPKK_00293_fwd", "CW1060_LLPMBPKK_00293_fwd"),
        ("###", "unnamed_oligo"),
    ],
)
def test_sanitize_idt_oligo_name(raw, expected):
    assert sanitize_idt_oligo_name(raw) == expected








def test_idt_order_df_to_idt_plates_preserves_row_order():
    idt_df = pd.DataFrame(
        {
            "Name": ["sgRNA_1", "sgRNA_2", "sgRNA_3"],
            "Sequence": ["ATGC", "GGCC", "TTAA"],
            "Concentration": ["25nm"] * 3,
            "Purification": ["STD"] * 3,
        }
    )

    plate = idt_order_df_to_idt_plates(idt_df)[0]

    assert list(plate.columns) == IDT_PLATE_COLUMNS
    assert plate["Name"].tolist() == ["sgRNA_1", "sgRNA_2", "sgRNA_3"]
    assert plate["Well Position"].tolist() == ["A1", "A2", "A3"]


def test_idt_order_df_to_idt_plates_reports_missing_columns():
    with pytest.raises(KeyError, match="Sequence"):
        idt_order_df_to_idt_plates(pd.DataFrame({"Name": ["a"]}))


def test_idt_plates_to_xlsx_bytes_writes_one_worksheet_per_plate():
    names = [f"oligo_{i}" for i in range(100)]
    plates = assign_plate_wells(names, ["ATGC"] * 100)

    payload = idt_plates_to_xlsx_bytes(plates)

    workbook = load_workbook(io.BytesIO(payload))
    assert workbook.sheetnames == ["Plate1", "Plate2"]

    first = workbook["Plate1"]
    assert [c.value for c in first[1]] == IDT_PLATE_COLUMNS
    assert first["A2"].value == "A1"
    assert first["B2"].value == "oligo_0"
    assert first["C2"].value == "ATGC"
    # 96 oligos plus the header row.
    assert first.max_row == 97

    assert workbook["Plate2"].max_row == 5


def test_idt_plates_to_xlsx_bytes_rejects_an_empty_order():
    with pytest.raises(ValueError, match="at least one plate"):
        idt_plates_to_xlsx_bytes([])


def test_plate_layout_matches_idts_sample_template_shape():
    """The three headers and row-major wells of IDT's own example file."""
    plate = assign_plate_wells(
        [
            "CW1060_LLPMBPKK_00293_fwd",
            "CW1061_LLPMBPKK_00293_rev",
            "CW1062_LLPMBPKK_00474_fwd",
        ],
        [
            "GGCGAGCAACGGAGGTACGGACAGGATGGACACCGAGAGTCAGG",
            "CGCAAGCCGCCACTCGAACGGAAGGTCACCGGTGGCGGCT",
            "GGCGAGCAACGGAGGTACGGACAGGATGCGGCTGGTCCAC",
        ],
    )[0]

    assert list(plate.columns) == ["Well Position", "Name", "Sequence"]
    assert plate["Well Position"].tolist() == ["A1", "A2", "A3"]
    assert plate["Name"].iloc[0] == "CW1060_LLPMBPKK_00293_fwd"
    assert plate["Sequence"].iloc[1].startswith("CGCAAGCCGCCACTCGAACGG")
